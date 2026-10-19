"""
Checklist reader — parses the Excel audit checklist into ChecklistItem objects.
"""

from pathlib import Path
from typing import List

from ai_project_auditor.models.checklist import ChecklistItem
from ai_project_auditor.core.exceptions import ChecklistNotFoundError


class ChecklistReader:
    """Reads and parses the Excel audit checklist."""

    SHEET_NAME = "Audit Checklist"

    # Column name mapping (flexible — matches by lowercased header)
    COL_CATEGORY = "category"
    COL_CHECK_ID = "check id"
    COL_CHECK = "check"
    COL_SEVERITY = "severity"
    COL_SCOPE = "scope"
    COL_TECHNOLOGY = "technology"
    COL_AI_CHECK = "ai check"
    COL_METHOD = "method"
    COL_DESCRIPTION = "description"
    COL_EXPECTED = "expected result"

    def __init__(self, checklist_path: str) -> None:
        self.checklist_path = Path(checklist_path)

    def load(self) -> List[ChecklistItem]:
        """Load and return all checklist items from the Excel file."""
        if not self.checklist_path.exists():
            raise ChecklistNotFoundError(
                f"Checklist not found: {self.checklist_path}"
            )

        try:
            import openpyxl
            wb = openpyxl.load_workbook(self.checklist_path, read_only=True, data_only=True)
        except Exception as exc:
            raise ChecklistNotFoundError(f"Cannot open checklist: {exc}") from exc

        if self.SHEET_NAME not in wb.sheetnames:
            raise ChecklistNotFoundError(
                f"Sheet '{self.SHEET_NAME}' not found in {self.checklist_path}. "
                f"Available sheets: {wb.sheetnames}"
            )

        ws = wb[self.SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        # Build header → column-index map
        header_row = [str(h).lower().strip() if h else "" for h in rows[0]]
        col_idx = {name: header_row.index(name) for name in header_row if name}

        items: List[ChecklistItem] = []
        for row in rows[1:]:
            if not any(row):
                continue  # skip blank rows

            def get(col_name: str, default="") -> str:
                idx = col_idx.get(col_name)
                if idx is None:
                    return default
                val = row[idx]
                return str(val).strip() if val is not None else default

            check_id = get(self.COL_CHECK_ID)
            if not check_id or check_id.lower() in ("check id", ""):
                continue

            ai_raw = get(self.COL_AI_CHECK, "No")
            ai_check = ai_raw.lower() in ("yes", "true", "1", "y")

            items.append(
                ChecklistItem(
                    check_id=check_id,
                    category=get(self.COL_CATEGORY),
                    check=get(self.COL_CHECK),
                    severity=get(self.COL_SEVERITY, "Low"),
                    scope=get(self.COL_SCOPE, "All"),
                    technology=get(self.COL_TECHNOLOGY, "All"),
                    ai_check=ai_check,
                    method=get(self.COL_METHOD, "Static"),
                    description=get(self.COL_DESCRIPTION),
                    expected_result=get(self.COL_EXPECTED),
                )
            )

        wb.close()
        return items
