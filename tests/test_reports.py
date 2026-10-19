"""
Tests for all Report Generators (JSON, Markdown, Excel, PDF).
"""

import os
import json
import tempfile
import openpyxl
from pathlib import Path
from ai_project_auditor.models.report import AuditReport, AuditStatistics, RiskSummary
from ai_project_auditor.models.finding import Finding, FindingSeverity, FindingConfidence, DetectionMethod
from ai_project_auditor.models.technology import TechnologyProfile
from ai_project_auditor.reports.json_report import JSONReportGenerator
from ai_project_auditor.reports.markdown_report import MarkdownReportGenerator
from ai_project_auditor.reports.excel_report import ExcelReportGenerator
from ai_project_auditor.reports.pdf_report import PDFReportGenerator


def _sample_report() -> AuditReport:
    tech = TechnologyProfile(python=True, fastapi=True, docker=True, primary_language="Python")
    finding = Finding(
        finding_id="FIND-001",
        check_id="SEC-001",
        title="Hardcoded API Key",
        category="Security",
        severity=FindingSeverity.CRITICAL,
        confidence=FindingConfidence.HIGH,
        file="config.py",
        line=10,
        evidence="API_KEY = '[REDACTED]'",
        description="Credentials hardcoded in config.",
        impact="Unauthorized access to APIs.",
        recommendation="Use environment variables.",
        ai_remediation="Change to os.getenv('API_KEY').",
        technology="Python",
        detection_method=DetectionMethod.STATIC,
        verified=True,
    )
    risk = RiskSummary.calculate([finding])
    stats = AuditStatistics(
        total_files=10,
        scanned_files=8,
        total_lines_of_code=1200,
        duration_seconds=1.45,
        timestamp="2026-09-28 10:00:00",
    )
    return AuditReport(
        project_name="TestApp",
        project_path="/test/path",
        technology_profile=tech,
        statistics=stats,
        risk_summary=risk,
        findings=[finding],
        executed_check_ids=["SEC-001", "OPS-002"],
        skipped_check_ids=["AI-002"],
    )


def test_json_report_generation():
    with tempfile.TemporaryDirectory() as tmp_dir:
        report = _sample_report()
        out_p = os.path.join(tmp_dir, "report.json")
        res_p = JSONReportGenerator.generate(report, out_p)

        assert os.path.exists(res_p)
        with open(res_p, "r", encoding="utf-8") as f:
            data = json.load(f)
        assert data["project_name"] == "TestApp"
        assert len(data["findings"]) == 1
        assert data["findings"][0]["check_id"] == "SEC-001"


def test_markdown_report_generation():
    with tempfile.TemporaryDirectory() as tmp_dir:
        report = _sample_report()
        out_p = os.path.join(tmp_dir, "report.md")
        res_p = MarkdownReportGenerator.generate(report, out_p)

        assert os.path.exists(res_p)
        content = Path(res_p).read_text(encoding="utf-8")
        assert "Audit Report: TestApp" in content
        assert "Executive Summary" in content
        assert "SEC-001" in content
        assert "Hardcoded API Key" in content


def test_excel_report_generation():
    with tempfile.TemporaryDirectory() as tmp_dir:
        report = _sample_report()
        out_p = os.path.join(tmp_dir, "report.xlsx")
        res_p = ExcelReportGenerator.generate(report, out_p)

        assert os.path.exists(res_p)
        wb = openpyxl.load_workbook(res_p)
        assert "Executive Summary" in wb.sheetnames
        assert "Findings" in wb.sheetnames
        assert "Checklist Coverage" in wb.sheetnames

        ws = wb["Findings"]
        assert ws.cell(2, 2).value == "SEC-001"


def test_pdf_report_generation():
    with tempfile.TemporaryDirectory() as tmp_dir:
        report = _sample_report()
        out_p = os.path.join(tmp_dir, "report.pdf")
        res_p = PDFReportGenerator.generate(report, out_p)

        assert os.path.exists(res_p)
        assert os.path.getsize(res_p) > 500  # Non-empty PDF


def test_finding_remediation_is_consistent_in_json_markdown_and_excel():
    with tempfile.TemporaryDirectory() as tmp_dir:
        report = _sample_report()
        guidance = report.findings[0].ai_remediation
        json_path = JSONReportGenerator.generate(report, os.path.join(tmp_dir, "report.json"))
        markdown_path = MarkdownReportGenerator.generate(report, os.path.join(tmp_dir, "report.md"))
        excel_path = ExcelReportGenerator.generate(report, os.path.join(tmp_dir, "report.xlsx"))

        assert json.loads(Path(json_path).read_text(encoding="utf-8"))["findings"][0]["ai_remediation"] == guidance
        assert guidance in Path(markdown_path).read_text(encoding="utf-8")
        workbook = openpyxl.load_workbook(excel_path)
        assert workbook["Findings"].cell(2, 13).value == guidance
